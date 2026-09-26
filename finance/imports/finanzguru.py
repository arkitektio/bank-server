"""Reading a Finanzguru transaction export.

Finanzguru exports every booking of every linked account into one file — officially ``.xlsx``;
the same columns saved as CSV (``;`` or tab separated, German numbers) are read too. One row per
booking, with its account (``Referenzkonto``), the bank fields, and Finanzguru's analysis:
category (Haupt-/Unterkategorie), whether it moved money between own accounts (``Umbuchung``),
contract, tags, a note.

What is taken, and where it goes (the rest stays in ``import_raw``):

* ``Buchungs-ID`` → ``import_ref`` ``fg:<id>``: the booking's identity across re-imports;
* ``Buchungstag``, ``Betrag`` (signed), ``Waehrung``, ``Beguenstigter/Auftraggeber``,
  ``Verwendungszweck`` → the bank fields; ``IBAN Beguenstigter/Auftraggeber`` only if it is an
  IBAN (Finanzguru puts PayPal e-mail addresses there too);
* ``Analyse-Haupt-/Unterkategorie`` → a category through the organization's mapping;
* ``Analyse-Umbuchung`` → a transfer between own accounts; ``Notiz`` → the note;
* ``Kontostand`` → the account's end-of-day balance.

``E-Ref`` is the SEPA end-to-end id, not a bank's entry reference, so it is kept in ``import_raw``
and never used as a fingerprint.

**Splits.** Finanzguru splits one booking into parts that carry the original's id in
``Referenz-Original-ID``. A booking must stay one row — the bank knows only the original, and
two rows would count it twice — so the parts are folded into their original (kept in
``import_raw["splits"]``); when only the parts are in the file, they are merged into one row
under the original's id.
"""

import csv
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from finance.models import normalize_iban

COLUMNS = [
    "Buchungstag",
    "Referenzkonto",
    "Name Referenzkonto",
    "Betrag",
    "Kontostand",
    "Waehrung",
    "Beguenstigter/Auftraggeber",
    "IBAN Beguenstigter/Auftraggeber",
    "Verwendungszweck",
    "E-Ref",
    "Mandatsreferenz",
    "Glaeubiger-ID",
    "Analyse-Hauptkategorie",
    "Analyse-Unterkategorie",
    "Analyse-Vertrag",
    "Analyse-Vertragsturnus",
    "Analyse-Vertrags-ID",
    "Analyse-Umbuchung",
    "Analyse-Vom frei verfuegbaren Einkommen ausgeschlossen",
    "Analyse-Umsatzart",
    "Analyse-Betrag",
    "Analyse-Woche",
    "Analyse-Monat",
    "Analyse-Quartal",
    "Analyse-Jahr",
    "Buchungs-ID",
    "Referenz-Original-ID",
    "Split-Typ",
    "Tags",
    "Notiz",
]
REQUIRED = ["Buchungstag", "Referenzkonto", "Betrag", "Waehrung", "Buchungs-ID"]

# Kept verbatim in ``import_raw`` (under these keys) — what the import knew beyond the bank fields.
EXTRA = {
    "E-Ref": "end_to_end_id",
    "Mandatsreferenz": "mandate_reference",
    "Glaeubiger-ID": "creditor_id",
    "Analyse-Vertrag": "contract",
    "Analyse-Vertragsturnus": "contract_interval",
    "Analyse-Vertrags-ID": "contract_id",
    "Analyse-Vom frei verfuegbaren Einkommen ausgeschlossen": "excluded_from_disposable_income",
    "Analyse-Umsatzart": "kind",
    "Analyse-Betrag": "analysis_amount",
    "Split-Typ": "split_type",
    "Tags": "tags",
}

TRUE = {"ja", "true", "wahr", "yes", "1", "x", "y", "j"}
_IBAN = re.compile(r"^[A-Z]{2}\d{2}[A-Z0-9]{10,30}$")
_THOUSANDS_ONLY = re.compile(r"^-?\d{1,3}(\.\d{3})+$")


class ExportError(ValueError):
    """The file is not a Finanzguru export this importer can read."""


@dataclass
class FgRow:
    """One booking of the export, as model-ready values."""

    booking_id: str
    reference: str
    reference_name: str | None
    booking_date: date
    amount: Decimal
    currency: str
    balance: Decimal | None
    counterparty: str | None
    counterparty_iban: str | None
    remittance: str | None
    main: str
    sub: str
    transfer: bool
    note: str | None
    original_id: str | None
    extra: dict[str, Any] = field(default_factory=dict)
    line: int = 0

    @property
    def import_ref(self) -> str:
        return f"fg:{self.booking_id}"


@dataclass
class ParsedExport:
    """Every booking of a file (splits already folded), and what was odd about it."""

    rows: list[FgRow]
    warnings: list[str] = field(default_factory=list)
    splits_folded: int = 0


def is_iban(value: str | None) -> bool:
    """True for something shaped like an IBAN (not an e-mail address, card number or account name)."""
    normalized = normalize_iban(value)
    return bool(normalized and _IBAN.match(normalized))


def _text(value: Any) -> str | None:  # noqa: ANN401 - a cell
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _decimal(value: Any, what: str) -> Decimal | None:  # noqa: ANN401
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    text = str(value).strip().replace(" ", "").replace(" ", "").replace("€", "").replace("EUR", "")
    if not text:
        return None
    if "," in text and "." in text:
        # The later separator is the decimal one: 1.234,56 (German) or 1,234.56.
        text = text.replace(".", "").replace(",", ".") if text.rfind(",") > text.rfind(".") else text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    elif _THOUSANDS_ONLY.match(text):
        text = text.replace(".", "")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation as error:
        raise ExportError(f"{what}: {value!r} is not an amount.") from error


def _date(value: Any, what: str) -> date:  # noqa: ANN401
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if text:
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S"):
            try:
                return datetime.strptime(text[:19] if "T" in text else text, fmt).date()
            except ValueError:
                continue
    raise ExportError(f"{what}: {value!r} is not a date.")


def _bool(value: Any) -> bool:  # noqa: ANN401
    if isinstance(value, bool):
        return value
    text = _text(value)
    return bool(text) and text.casefold() in TRUE


def _json_safe(value: Any) -> Any:  # noqa: ANN401
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, float):
        return str(Decimal(str(value)))
    return value


def _table(content: bytes) -> list[list[Any]]:
    """The file's cells, row by row: an xlsx workbook's first sheet, else CSV."""
    if content[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise ExportError(f"Not a readable xlsx workbook: {error}") from error
        try:
            return [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
        finally:
            workbook.close()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ExportError("The file is neither xlsx nor text.")
    first = text.split("\n", 1)[0]
    delimiter = max([";", "\t", ","], key=first.count)
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _header(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    """Where the header row is (a title row may come first) and each column's index."""
    for index, row in enumerate(rows[:10]):
        names = {(_text(cell) or ""): position for position, cell in enumerate(row)}
        if all(name in names for name in REQUIRED):
            return index, names
    raise ExportError(f"No Finanzguru header found: the columns {', '.join(REQUIRED)} are required.")


def parse(content: bytes) -> ParsedExport:
    """Read an export into bookings, with splits folded; raises :class:`ExportError` if it cannot."""
    table = _table(content)
    start, columns = _header(table)
    warnings = []
    unknown = sorted(set(columns) - set(COLUMNS) - {""})
    missing = [name for name in COLUMNS if name not in columns]
    if unknown:
        warnings.append(f"Columns this importer does not know were kept in import_raw: {', '.join(unknown)}.")
    if missing:
        warnings.append(f"Columns missing from the file: {', '.join(missing)}.")

    def cell(row: list[Any], name: str) -> Any:  # noqa: ANN401
        position = columns.get(name)
        return row[position] if position is not None and position < len(row) else None

    rows: list[FgRow] = []
    seen: set[str] = set()
    for line, raw in enumerate(table[start + 1 :], start=start + 2):
        if not any(_text(value) for value in raw):
            continue
        where = f"line {line}"
        booking_id = _text(cell(raw, "Buchungs-ID"))
        reference = _text(cell(raw, "Referenzkonto"))
        amount = _decimal(cell(raw, "Betrag"), f"{where}, Betrag")
        if not booking_id or not reference or amount is None:
            warnings.append(f"{where}: skipped (no Buchungs-ID, Referenzkonto or Betrag).")
            continue
        if booking_id in seen:
            warnings.append(f"{where}: Buchungs-ID {booking_id} appears twice; the second is skipped.")
            continue
        seen.add(booking_id)
        party_iban = _text(cell(raw, "IBAN Beguenstigter/Auftraggeber"))
        extra = {key: _json_safe(value) for name, key in EXTRA.items() if (value := cell(raw, name)) not in (None, "")}
        if party_iban and not is_iban(party_iban):
            extra["counterparty_account"] = party_iban
        for name in unknown:
            if (value := cell(raw, name)) not in (None, ""):
                extra.setdefault("other", {})[name] = _json_safe(value)
        original = _text(cell(raw, "Referenz-Original-ID"))
        rows.append(
            FgRow(
                booking_id=booking_id,
                reference=reference,
                reference_name=_text(cell(raw, "Name Referenzkonto")),
                booking_date=_date(cell(raw, "Buchungstag"), f"{where}, Buchungstag"),
                amount=amount,
                currency=(_text(cell(raw, "Waehrung")) or "EUR").upper()[:3],
                balance=_decimal(cell(raw, "Kontostand"), f"{where}, Kontostand"),
                counterparty=_text(cell(raw, "Beguenstigter/Auftraggeber")),
                counterparty_iban=normalize_iban(party_iban) if is_iban(party_iban) else None,
                remittance=_text(cell(raw, "Verwendungszweck")),
                main=_text(cell(raw, "Analyse-Hauptkategorie")) or "",
                sub=_text(cell(raw, "Analyse-Unterkategorie")) or "",
                transfer=_bool(cell(raw, "Analyse-Umbuchung")),
                note=_text(cell(raw, "Notiz")),
                original_id=original if original and original != booking_id else None,
                extra=extra,
                line=line,
            )
        )
    folded, count = fold_splits(rows, warnings)
    return ParsedExport(rows=folded, warnings=warnings, splits_folded=count)


def _split_part(row: FgRow) -> dict[str, Any]:
    return {"booking_id": row.booking_id, "amount": str(row.amount), "main": row.main, "sub": row.sub, "note": row.note, "remittance": row.remittance}


def fold_splits(rows: list[FgRow], warnings: list[str]) -> tuple[list[FgRow], int]:
    """One row per real booking: split parts folded into their original (or merged when it is absent)."""
    parts: dict[str, list[FgRow]] = defaultdict(list)
    for row in rows:
        if row.original_id:
            parts[row.original_id].append(row)
    if not parts:
        return rows, 0
    by_id = {row.booking_id: row for row in rows}
    out: list[FgRow] = []
    for row in rows:
        if row.original_id:
            continue  # a part: folded below, into its original's place
        group = parts.get(row.booking_id)
        if group:
            out.append(_with_parts(row, group, warnings))
        else:
            out.append(row)
    for original_id, group in parts.items():
        if original_id in by_id:
            continue
        # Only the parts are in the file: they are one booking, under the original's id.
        first = group[0]
        merged = FgRow(**{**first.__dict__, "booking_id": original_id, "original_id": None, "amount": sum((p.amount for p in group), Decimal("0")), "extra": dict(first.extra)})
        out.append(_with_parts(merged, group, warnings, check_total=False))
    return out, sum(len(group) for group in parts.values())


def _with_parts(row: FgRow, group: list[FgRow], warnings: list[str], check_total: bool = True) -> FgRow:
    total = sum((p.amount for p in group), Decimal("0"))
    if check_total and total != row.amount:
        warnings.append(f"Buchungs-ID {row.booking_id}: its split parts sum to {total}, not {row.amount}; the booking keeps {row.amount}.")
    categories = {(p.main, p.sub) for p in group if p.main}
    if len(categories) == 1:
        row.main, row.sub = next(iter(categories))
    elif categories:
        largest = max(group, key=lambda p: abs(p.amount))
        row.main, row.sub = largest.main, largest.sub
    row.note = row.note or next((p.note for p in group if p.note), None)
    row.transfer = row.transfer or all(p.transfer for p in group)
    row.extra = {**row.extra, "splits": [_split_part(p) for p in group]}
    return row
